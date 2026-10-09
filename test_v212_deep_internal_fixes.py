import source_bundle
from pathlib import Path
import json, os, tempfile

ROOT = Path(__file__).resolve().parent
WEB = source_bundle.webapp_text()
CHAT = source_bundle.chat_view_text()
DB = (ROOT/'db.py').read_text(encoding='utf-8')
CORE = (ROOT/'analysis_core.py').read_text(encoding='utf-8')


def test_medication_delete_uses_csrf_credentials_and_checks_response():
    assert "async function deletePlan(id)" in WEB
    assert "headers:medMutationHeaders(false)" in WEB
    assert "credentials:'same-origin'" in WEB
    assert "if(!r.ok||!d.ok)throw new Error" in WEB
    assert "✓ تم حذف التذكير" in WEB
    assert "e.stopPropagation()" in WEB


def test_medication_status_mutations_are_csrf_protected_too():
    assert "headers:medMutationHeaders(true)" in WEB
    assert "/api/meds/log" in WEB
    assert "/api/meds/snooze" in WEB


def test_edit_answers_opens_real_prefilled_form():
    for token in (
        'answer-edit-panel', 'editAnswerAge', 'editAnswerGender',
        'editAnswerSymptoms', 'editAnswerDuration', 'editAnswerSeverity',
        'editAnswerConditions', 'editAnswerMeds', 'editAnswerAllergies',
        'editAnswerNotes', 'data-answer-edit-save', 'data-answer-edit-step',
    ):
        assert token in CHAT
    assert "Your current answers are pre-filled" in CHAT
    assert "حفظ وإعادة التحليل" in CHAT
    assert "runAnalysis();" in CHAT


def test_edit_answers_preserves_reanalysis_link_and_resets_stale_pattern_if_symptoms_change():
    assert "state.previous_record_id=id" in CHAT
    assert "oldSymptoms!==state.symptoms.join('||')" in CHAT
    assert "state.pattern_context_done=false" in CHAT
    assert "differentialAnswers=[]" in CHAT


def test_persisted_result_has_symptom_input_snapshot_for_export_recovery():
    assert '"input_snapshot": {' in CORE
    assert '"symptoms": list(d.get("symptoms") or [])' in CORE


def test_admin_export_can_recover_legacy_symptoms_from_result_json():
    block = DB[DB.index('def admin_symptom_trial_rows():'):DB.index('def admin_symptom_trial_summary():')]
    assert 'LEFT JOIN results' in block
    assert "snapshot.get('symptoms')" in block
    assert "data.get('symptoms')" in block
    assert "norm.get('canonical')" in block
    assert "norm.get('unmatched')" in block


def test_runtime_excel_contains_symptoms_and_legacy_recovery(monkeypatch):
    import db
    from openpyxl import load_workbook
    tmp = tempfile.NamedTemporaryFile(suffix='.db', delete=False)
    tmp.close()
    try:
        monkeypatch.setattr(db, 'DATABASE_URL', '')
        monkeypatch.setattr(db, 'USE_POSTGRES', False)
        monkeypatch.setattr(db, 'PH', '?')
        monkeypatch.setattr(db, 'DB_PATH', tmp.name)
        monkeypatch.setattr(db, '_DB_READY_KEY', None)
        db.init_db()
        db.save_analysis_record_with_result(
            'account-v212','ar',30,'m',['صداع','غثيان'],'1-3 أيام',2,'low',
            {'ok':True,'symptoms':['صداع','غثيان'],'input_snapshot':{'symptoms':['صداع','غثيان']}},
        )
        conn=db._conn(); c=conn.cursor()
        c.execute(
            "INSERT INTO records (user_hash,timestamp,lang,age,gender,symptoms,conditions,medications,duration,severity,urgency,member_id) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (db._hash_user('account-old'),'2026-10-03T00:00:00+00:00','ar',25,'f','','','','يوم',1,'low',0),
        )
        rid=c.lastrowid
        payload={'symptom_normalization':{'canonical':[{'original':'ألم بطن'}],'unmatched':['دوخة']}}
        c.execute("INSERT INTO results (user_hash,record_id,data) VALUES (?,?,?)",(db._hash_user('account-old'),rid,json.dumps(payload,ensure_ascii=False)))
        conn.commit(); conn.close()
        rows=db.admin_symptom_trial_rows()
        assert any('صداع' in x['symptoms'] and 'غثيان' in x['symptoms'] for x in rows)
        assert any('ألم بطن' in x['symptoms'] and 'دوخة' in x['symptoms'] for x in rows)
        wb=load_workbook(db.admin_export_symptom_trials_xlsx(), data_only=True)
        ws=wb['Symptom Trials']
        exported=[ws.cell(r,3).value or '' for r in range(5,ws.max_row+1)]
        assert any('صداع' in x and 'غثيان' in x for x in exported)
        assert any('ألم بطن' in x and 'دوخة' in x for x in exported)
    finally:
        try: os.unlink(tmp.name)
        except OSError: pass


def test_emergency_assessments_are_persisted_for_signed_in_users():
    assert 'Emergency symptom analysis could not be persisted' in WEB
    assert 'emergency_persisted["symptoms"] = list(patient.get("symptoms") or [])' in WEB
    assert 'result["record_saved"] = bool(emergency_record_id)' in WEB
    assert 'db.save_analysis_record_with_result(' in WEB
