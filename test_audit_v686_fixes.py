"""Reproduce the audit's data-loss/privacy/consent gaps with synthetic data."""
import io
import json
import shutil
import uuid
from pathlib import Path

import pytest

import db
import family_archive
import privacy_features
import research_study


def account():
    privacy_features.init_schema()
    uid, error = db.create_ss_user(f"audit-{uuid.uuid4().hex}@example.test", "Synthetic Owner", "SyntheticPassword123!")
    assert uid and not error
    return uid, f"account-{uid}"


def seed_health(uid, owner):
    member = db.save_member(owner, "mother", "Synthetic Family", 45, "female", conditions="asthma", notes="Private archived note")
    rid = db.save_record(owner, "en", 30, "female", ["headache"], "2 days", 3, "low")
    db.save_symptom_relief_log(owner, rid, "headache", "rest")
    db.save_blood_test(owner, {"indicators": [{"key": "hgb", "value": 13}]})
    db.save_vital(owner, 0, "weight", 60)
    db.save_web_daily_checkin(uid, 3, "2026-01-01")
    assert family_archive.delete_member(owner, uid, member, delete_data=False)
    return member, rid


@pytest.mark.parametrize("delete_account", [False, True])
def test_delete_removes_archives_relief_and_daily_tracking_but_not_other_owner(delete_account):
    uid, owner = account(); other_uid, other = account()
    member, _ = seed_health(uid, owner); other_member, _ = seed_health(other_uid, other)
    if delete_account:
        db.delete_ss_user(uid)
    else:
        privacy_features.delete_health_data(uid, owner)
        assert db.get_ss_user(uid)
    assert family_archive.list_archived(owner) == []
    assert not family_archive.restore_member(owner, member)
    assert db.get_symptom_relief_summary(owner, "headache") == []
    assert db.get_web_daily_checkin_history(uid) == []
    assert family_archive.list_archived(other)[0]["id"] == other_member
    assert db.get_symptom_relief_summary(other, "headache")
    assert db.get_web_daily_checkin_history(other_uid)


def test_owner_export_contains_all_health_types_and_does_not_truncate():
    uid, owner = account(); other_uid, other = account()
    member, rid = seed_health(uid, owner)
    other_member = db.save_member(other, "mother", "OTHER-OWNER-PRIVATE", 50, "female")
    conn = db._conn()
    try:
        cur = conn.cursor()
        cur.executemany(
            f"INSERT INTO records(user_hash,timestamp,lang,age,gender,symptoms) VALUES ({','.join([db.PH]*6)})",
            [(db._hash_user(owner), "2026-01-01", "en", 30, "female", "headache")] * 1001,
        )
        cur.executemany(
            f"INSERT INTO ss_chat_history(user_id,role,content,timestamp) VALUES ({','.join([db.PH]*4)})",
            [(uid, "user", f"Synthetic message {i}", "2026-01-01") for i in range(501)],
        )
        conn.commit()
    finally:
        conn.close()
    export = privacy_features._safe_user_export(uid, owner)
    assert export["complete"] is True
    assert len(export["analysis_history"]) == 1002
    assert len(export["assistant_history"]) == 501
    assert export["blood_tests"][0]["data"]["indicators"][0]["value"] == 13
    assert export["vitals"][0]["v1"] == 60
    assert export["archived_family_members"][0]["member_id"] == member
    assert export["symptom_relief_logs"][0]["record_id"] == rid
    assert export["web_daily_checkins"] == [{"date": "2026-01-01", "value": 3}]
    blob = json.dumps(export)
    assert "OTHER-OWNER-PRIVATE" not in blob and "password_hash" not in blob and "token_hash" not in blob
    with pytest.raises(ValueError, match="owner_mismatch"):
        privacy_features._safe_user_export(uid, other)


def test_guest_research_only_is_saved_and_withdrawal_does_not_change_other_purpose():
    import guest_trials
    import admin_complete
    owner = "guest-audit-" + uuid.uuid4().hex
    privacy_features.save_consent(owner, None, True, False, True)
    rid = guest_trials.save(owner, "en", {"age": 30, "symptoms": ["headache"]}, {}, False, False, True)
    assert rid
    conn = db._conn()
    try:
        cur = conn.cursor(); cur.execute(f"SELECT analytics_eligible,research_eligible FROM records WHERE id={db.PH}", (rid,))
        assert tuple(cur.fetchone()) == (0, 1)
    finally:
        conn.close()
    privacy_features.withdraw_research(owner)
    conn = db._conn()
    try:
        cur = conn.cursor(); cur.execute(f"SELECT research_eligible FROM records WHERE id={db.PH}", (rid,)); assert cur.fetchone()[0] == 0
    finally:
        conn.close()
    assert guest_trials.save(owner, "en", {"age": 30, "symptoms": ["headache"]}, {}, False, False, False) is None


def test_research_json_and_excel_never_export_patient_prose_or_echoed_identifiers():
    from openpyxl import load_workbook
    import admin_complete
    uid, owner = account()
    pii = "UnknownPrivateName synthetic-person@example.test +966500001234"
    result = {key: pii for key in ("possible_conditions", "why_result", "danger_signs", "when_to_seek_care", "home_care", "risk_label", "knowledge_last_updated")}
    result.update({"knowledge_matches": [{"slug": pii, "name_en": pii}], "data_quality": {"score": 80, "extra": pii},
                   "medical_sources": [{"url": "https://example.test/" + pii}], "recommendations": [{"title": pii, "tip": pii}], "confidence": pii})
    rid = db.save_analysis_record_with_result(owner, "en", 30, "female", ["headache", pii], pii, 3, "low", result, conditions=pii, medications=pii)
    privacy_features.set_record_consent_eligibility(rid, False, True)
    payload = admin_complete.admin_analysis_results()
    blob = json.dumps(payload)
    assert "UnknownPrivateName" not in blob and "synthetic-person" not in blob and "+966500001234" not in blob
    assert any("headache" in " ".join(r["symptoms"]).lower() for r in payload["records"])
    workbook = load_workbook(admin_complete.export_admin_workbook("pilot"))
    cells = " ".join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
    assert "UnknownPrivateName" not in cells and "synthetic-person" not in cells and "+966500001234" not in cells
    assert "Structured catalog terms only" in cells
    assert pii in json.dumps(privacy_features._safe_user_export(uid, owner))  # owner's own data stays intact


def test_research_signature_detects_changed_added_and_removed_request_and_kb_files(tmp_path, monkeypatch):
    original = research_study.BASE_DIR
    for rel in research_study.CORE_FILES:
        src = original / rel
        if src.exists():
            dest = tmp_path / rel; dest.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(src, dest)
    routes = tmp_path / "routes"; routes.mkdir()
    path = routes / "assistant.py"; path.write_text("original")
    monkeypatch.setattr(research_study, "BASE_DIR", tmp_path)
    before = research_study.algorithm_signature()["sha256"]
    path.write_text("changed")
    changed = research_study.algorithm_signature()["sha256"]; assert changed != before
    path.unlink(); assert research_study.algorithm_signature()["sha256"] != changed
    baseline = research_study.algorithm_signature()["sha256"]
    (tmp_path / "medical_knowledge_expansion_new.py").write_text("new knowledge")
    assert research_study.algorithm_signature()["sha256"] != baseline


def test_welcome_preview_card_is_not_rotated_on_phones():
    """The decorative card behind the language chooser used to be tilted -5deg, so its text looked slanted on phones."""
    import welcome_phone
    assert "rotate(" not in welcome_phone.CSS and "translate(-50%,-50%)!important" in welcome_phone.CSS


def test_trial_export_excludes_identifiers_and_formulas_through_analysis_api():
    import webapp
    from openpyxl import load_workbook
    client = webapp.app.test_client()
    assert client.post('/api/consent/preferences', json={
        'service_usage': True, 'analytics_research': True, 'research_participation': False,
    }).status_code == 200
    marker = 'PRIVATE-SYNTHETIC-PERSON demo@example.test'
    response = client.post('/api/analyze', json={
        'age': 30, 'gender': 'female', 'symptoms': ['headache', marker],
        'duration': '=1+1', 'severity': 3, 'lang': 'en', 'notes': '',
    })
    assert response.status_code == 200
    workbook = load_workbook(db.admin_export_symptom_trials_xlsx(), data_only=False)
    cells = [cell for row in workbook.active for cell in row]
    rendered = ' '.join(str(c.value) for c in cells)
    assert 'PRIVATE-SYNTHETIC-PERSON' not in rendered and 'demo@example.test' not in rendered
    assert 'headache' in rendered.lower()
    assert not any(c.data_type == 'f' for c in cells)
    conn = db._conn()
    try:
        cur = conn.cursor(); cur.execute('SELECT symptoms,duration FROM records ORDER BY id DESC LIMIT 1')
        stored = cur.fetchone()
        assert marker in stored[0] and stored[1] == '=1+1'
    finally:
        conn.close()


def test_trial_export_legacy_recovery_also_excludes_private_text():
    owner = 'legacy-private-' + uuid.uuid4().hex
    rid = db.save_analysis_record_with_result(owner, 'en', 30, 'female', [], '2 days PRIVATE-NAME', 3, 'low',
        {'input_snapshot': {'symptoms': ['headache', 'PRIVATE-NAME email@example.test']}})
    privacy_features.set_record_consent_eligibility(rid, True, False)
    rows = db.admin_symptom_trial_rows()
    rendered = json.dumps(rows)
    assert 'PRIVATE-NAME' not in rendered and 'email@example.test' not in rendered
    assert any(row['duration'] == 'Days' for row in rows)


@pytest.mark.parametrize('payload', ['=1+1', '+1+1', '-1+1', '@SUM(1,1)', '\t=1+1'])
def test_trial_workbook_keeps_untrusted_strings_as_text(monkeypatch, payload):
    from openpyxl import load_workbook
    monkeypatch.setattr(db, 'admin_symptom_trial_rows', lambda: [
        {'age': 30, 'gender': payload, 'symptoms': payload, 'duration': payload}])
    ws = load_workbook(db.admin_export_symptom_trials_xlsx(), data_only=False).active
    for col in ('B', 'C', 'D'):
        assert ws[f'{col}5'].data_type == 's'
        assert ws[f'{col}5'].value == payload
