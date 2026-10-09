"""Privacy-preserving operational analytics for SymptoSense Admin.

This module intentionally exposes only aggregate/anonymous information. It never
returns names, emails, phones, authentication secrets, raw health profiles, or
private conversations.
"""
from __future__ import annotations
import logging

import io
import os
import re
import secrets
import hashlib
import threading
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import db
import platform_v2
import privacy_features

PH = db.PH
_EPHEMERAL_ANALYTICS_SECRET = secrets.token_hex(32)
_SCHEMA_READY_KEY = None
_SCHEMA_LOCK = threading.Lock()

JOURNEY_STAGES = ("home", "start_analysis", "symptoms", "questionnaire", "analysis", "result", "reanalyze", "report", "quick_start", "full_start")
LIVE_TYPES = {
    "analysis_started": ("Analysis started", "analysis"),
    "analysis_complete": ("Analysis completed", "analysis"),
    "assistant_use": ("Smart Assistant used", "assistant"),
    "report_generated": ("Report generated", "reports"),
    "reanalysis_started": ("Re-analysis started", "analysis"),
    "new_account": ("New account created", "authentication"),
    "source_accessed": ("Medical source accessed", "system"),
    "error": ("System request error", "system"),
    "safety_alert": ("Safety alert triggered", "system"),
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _serial() -> str:
    return "SERIAL PRIMARY KEY" if db.USE_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"


def _columns(cur, table: str) -> set[str]:
    if db.USE_POSTGRES:
        cur.execute("SELECT column_name FROM information_schema.columns WHERE table_name=%s", (table,))
        return {r[0] for r in cur.fetchall()}
    cur.execute(f"PRAGMA table_info({table})")
    return {r[1] for r in cur.fetchall()}


def init_schema() -> None:
    """Initialize operational analytics tables once per database identity."""
    global _SCHEMA_READY_KEY
    db.init_db()
    key = db._database_identity()
    if _SCHEMA_READY_KEY == key:
        return
    with _SCHEMA_LOCK:
        if _SCHEMA_READY_KEY == key:
            return
        conn = db._conn(); c = conn.cursor()
        try:
            c.execute(f"""
                CREATE TABLE IF NOT EXISTS ss_journey_events (
                    id {_serial()}, session_hash TEXT NOT NULL, stage TEXT NOT NULL,
                    device_type TEXT NOT NULL, created_at TEXT NOT NULL
                )
            """)
            c.execute("CREATE INDEX IF NOT EXISTS idx_ss_journey_time ON ss_journey_events(created_at)")
            c.execute("CREATE INDEX IF NOT EXISTS idx_ss_journey_session ON ss_journey_events(session_hash, created_at)")
            c.execute("""
                CREATE TABLE IF NOT EXISTS ss_session_activity (
                    session_hash TEXT PRIMARY KEY, device_type TEXT NOT NULL,
                    last_seen TEXT NOT NULL
                )
            """)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
        _SCHEMA_READY_KEY = key


def session_hash(raw_session_id: str) -> str:
    salt = os.environ.get("ANALYTICS_SESSION_SALT") or os.environ.get("WEB_SECRET") or _EPHEMERAL_ANALYTICS_SECRET
    return hashlib.sha256((salt + "|" + str(raw_session_id or "")).encode("utf-8")).hexdigest()[:32]


def touch_session(raw_session_id: str, user_agent: str) -> None:
    if not raw_session_id:
        return
    init_schema(); sh = session_hash(raw_session_id); dev = platform_v2.device_type(user_agent); now = _now()
    conn = db._conn(); c = conn.cursor()
    try:
        if db.USE_POSTGRES:
            c.execute("INSERT INTO ss_session_activity(session_hash,device_type,last_seen) VALUES(%s,%s,%s) ON CONFLICT(session_hash) DO UPDATE SET device_type=EXCLUDED.device_type,last_seen=EXCLUDED.last_seen", (sh, dev, now))
        else:
            c.execute("INSERT INTO ss_session_activity(session_hash,device_type,last_seen) VALUES(?,?,?) ON CONFLICT(session_hash) DO UPDATE SET device_type=excluded.device_type,last_seen=excluded.last_seen", (sh, dev, now))
        conn.commit()
    finally:
        conn.close()


def record_journey(raw_session_id: str, stage: str, user_agent: str) -> bool:
    if not raw_session_id or stage not in JOURNEY_STAGES:
        return False
    init_schema(); sh = session_hash(raw_session_id); dev = platform_v2.device_type(user_agent)
    conn = db._conn(); c = conn.cursor()
    try:
        # Prevent noisy duplicate beacons while keeping repeated flows possible.
        since = (datetime.now(timezone.utc) - timedelta(seconds=20)).isoformat()
        c.execute(f"SELECT 1 FROM ss_journey_events WHERE session_hash={PH} AND stage={PH} AND created_at>={PH} LIMIT 1", (sh, stage, since))
        if c.fetchone():
            return False
        c.execute(f"INSERT INTO ss_journey_events(session_hash,stage,device_type,created_at) VALUES({PH},{PH},{PH},{PH})", (sh, stage, dev, _now()))
        conn.commit(); return True
    finally:
        conn.close()


def mode_completion(days: int = 30) -> dict:
    """Completion rate of the quick vs. full path, from consented, anonymous journey beacons.

    A session "completes" when it records a ``result`` stage at or after its first
    ``quick_start`` / ``full_start`` beacon.
    """
    since = (datetime.now(timezone.utc) - timedelta(days=max(1, min(int(days), 365)))).isoformat()
    init_schema()
    conn = db._conn()
    try:
        c = conn.cursor()
        c.execute(f"SELECT session_hash, stage, created_at FROM ss_journey_events WHERE created_at>={PH} ORDER BY created_at", (since,))
        rows = c.fetchall()
    finally:
        conn.close()
    first_start, result_at = {}, defaultdict(list)
    for sid, stage, ts in rows:
        sid, stage, ts = str(sid), str(stage), str(ts)
        if stage in ("quick_start", "full_start"):
            first_start.setdefault((sid, stage), ts)
        elif stage == "result":
            result_at[sid].append(ts)
    out = {}
    for mode in ("quick_start", "full_start"):
        started = [(sid, ts) for (sid, st), ts in first_start.items() if st == mode]
        done = sum(1 for sid, ts in started if any(r >= ts for r in result_at.get(sid, ())))
        out[mode.replace("_start", "")] = {
            "started": len(started), "completed": done,
            "completion_rate": round(100.0 * done / len(started), 1) if started else None,
        }
    return {"days": days, **out}


def _bounds(period: str = "30d", start: str | None = None, end: str | None = None):
    now = datetime.now(timezone.utc)
    period = (period or "30d").lower()
    if period == "today":
        local = datetime.now(ZoneInfo("Asia/Riyadh")); s_local = local.replace(hour=0, minute=0, second=0, microsecond=0)
        return s_local.astimezone(timezone.utc), now
    if period == "7d": return now - timedelta(days=7), now
    if period == "30d": return now - timedelta(days=30), now
    if period == "90d": return now - timedelta(days=90), now
    if period == "180d": return now - timedelta(days=180), now
    if period == "all": return None, now
    if period == "custom" and start:
        try:
            s = datetime.fromisoformat(start[:10] + "T00:00:00+00:00")
            e = datetime.fromisoformat((end or start)[:10] + "T23:59:59+00:00")
            if e < s: s, e = e, s
            return s, min(e, now)
        except Exception:
            logging.getLogger(__name__).debug("Handled exception in _bounds; fallback applied (handler 131)")
            pass
    return now - timedelta(days=30), now


def dropoff_analysis(period="30d", start=None, end=None, device="") -> dict:
    init_schema(); since, until = _bounds(period, start, end)
    conn = db._conn(); c = conn.cursor()
    try:
        clauses = [f"created_at<={PH}"]; params = [until.isoformat()]
        if since is not None: clauses.append(f"created_at>={PH}"); params.append(since.isoformat())
        if device in {"mobile", "desktop", "tablet"}: clauses.append(f"device_type={PH}"); params.append(device)
        c.execute("SELECT session_hash,stage,created_at FROM ss_journey_events WHERE " + " AND ".join(clauses) + " ORDER BY created_at", tuple(params))
        events = c.fetchall()
    finally:
        conn.close()
    by_session: dict[str, set[str]] = defaultdict(set)
    for sid, stage, _ in events:
        by_session[str(sid)].add(str(stage))
    core = ["home", "start_analysis", "symptoms", "questionnaire", "analysis", "result"]
    labels = {"home":"Home","start_analysis":"Start Analysis","symptoms":"Symptoms","questionnaire":"Questionnaire","analysis":"Analysis","result":"Result"}
    counts = {s: sum(1 for st in by_session.values() if s in st) for s in core}
    rows=[]; prev=None; highest=None
    for s in core:
        n=counts[s]
        conversion = 100.0 if prev is None or prev == 0 else min(100.0, round(n*100.0/prev,1))
        drop = 0.0 if prev is None or prev == 0 else max(0.0, round((prev-n)*100.0/prev,1))
        row={"stage":s,"label":labels[s],"users":n,"conversion_rate":conversion,"dropoff_rate":drop}
        rows.append(row)
        if prev is not None and (highest is None or drop > highest["dropoff_rate"]): highest=row
        prev=n
    base=counts["home"] or counts["start_analysis"] or 0
    complete=counts["result"]
    completion=round(complete*100.0/base,1) if base else 0.0
    return {"stages":rows,"completion_rate":completion,"highest_dropoff":highest if base else None,"sessions":len(by_session),"period":period,"device":device or "all"}


def live_activity(limit=80, category="all", window="hour") -> dict:
    init_schema(); now=datetime.now(timezone.utc)
    if window == "7d": since=now-timedelta(days=7)
    elif window == "today": since=_bounds("today")[0]
    else: since=now-timedelta(hours=1)
    conn=db._conn(); c=conn.cursor()
    try:
        c.execute(f"SELECT event_type,response_status,created_at FROM ss_usage_events WHERE created_at>={PH} ORDER BY created_at DESC LIMIT {PH}", (since.isoformat(), max(1,min(300,int(limit)))))
        raw=c.fetchall()
        active_since=(now-timedelta(minutes=5)).isoformat()
        c.execute(f"SELECT COUNT(*) FROM ss_session_activity WHERE last_seen>={PH}", (active_since,)); active=int(c.fetchone()[0] or 0)
        today_start=_bounds("today")[0].isoformat()
        c.execute(f"SELECT COUNT(*) FROM records WHERE timestamp>={PH}", (today_start,)); analyses_today=int(c.fetchone()[0] or 0)
        c.execute(f"SELECT COUNT(*) FROM ss_usage_events WHERE event_type='report_generated' AND created_at>={PH}", (today_start,)); reports_today=int(c.fetchone()[0] or 0)
        progress_since=(now-timedelta(hours=2)).isoformat()
        c.execute(f"SELECT session_hash,stage FROM ss_journey_events WHERE created_at>={PH}", (progress_since,)); jr=c.fetchall()
    finally: conn.close()
    stage_map=defaultdict(set)
    for sid,st in jr: stage_map[str(sid)].add(str(st))
    in_progress=sum(1 for st in stage_map.values() if "analysis" in st and "result" not in st)
    events=[]; buckets=Counter()
    for typ,status,ts in raw:
        meta=LIVE_TYPES.get(str(typ))
        if not meta: continue
        label,cat=meta
        if category not in ("", "all") and cat != category: continue
        events.append({"event":typ,"label":label,"category":cat,"status":"error" if int(status or 0)>=400 else "normal","timestamp":ts})
        buckets[str(ts)[:13]+":00"] += 1
    trend=[{"date":k,"count":v} for k,v in sorted(buckets.items())]
    return {"metrics":{"active_sessions":active,"analyses_today":analyses_today,"analyses_in_progress":in_progress,"reports_today":reports_today},"events":events,"trend":trend,"window":window}


# ---------------- Medication analytics export ----------------
def _age_group(age) -> str:
    try: a=int(age)
    except Exception: logging.getLogger(__name__).debug("Handled exception in _age_group; fallback applied (handler 203)"); return "Unknown"
    if a < 18: return "Under 18"
    if a <= 25: return "18–25"
    if a <= 35: return "26–35"
    if a <= 45: return "36–45"
    if a <= 55: return "46–55"
    return "56+"


def _gender(v) -> str:
    s=str(v or "").strip().lower()
    if s in {"m","male","ذكر","man"}: return "Male"
    if s in {"f","female","أنثى","انثى","woman"}: return "Female"
    return "Unknown"


def _risk(v) -> str:
    s=str(v or "").strip().lower()
    if s in {"high","urgent","emergency"}: return "Urgent"
    if s in {"medium","review","needs_followup","needs follow-up","today"}: return "Needs Follow-up"
    if s in {"low","normal"}: return "Low Risk"
    return "Unknown"


def _split_items(value: str) -> list[str]:
    bad={"none","no","n/a","na","لا","لا يوجد","لايوجد","بدون","nothing","-"}
    out=[]
    for x in re.split(r"[,،;\n|]+", str(value or "")):
        x=re.sub(r"\s+", " ", x).strip(" .-")
        x=re.sub(r"\b\d+(?:\.\d+)?\s*(?:mg|mcg|g|ml|iu|units?)\b", "", x, flags=re.I).strip()
        if x and x.lower() not in bad and len(x)<=100: out.append(x)
    return out


def _norm(v: str) -> str:
    return re.sub(r"\s+"," ",str(v or "").strip().casefold())


def _records_for_filters(filters: dict) -> list[dict]:
    privacy_features.init_schema()
    period=filters.get("period") or "30d"; since,until=_bounds(period,filters.get("start"),filters.get("end"))
    conn=db._conn(); c=conn.cursor()
    try:
        clauses=["analytics_eligible=1",f"timestamp<={PH}"]; params=[until.isoformat()]
        if since is not None: clauses.append(f"timestamp>={PH}");params.append(since.isoformat())
        c.execute("SELECT id,user_hash,timestamp,age,gender,symptoms,medications,duration,severity,urgency FROM records WHERE "+" AND ".join(clauses)+" ORDER BY timestamp",tuple(params))
        rows=[]
        for r in c.fetchall():
            item={"id":r[0],"user_hash":str(r[1]),"timestamp":str(r[2] or ""),"age":r[3],"age_group":_age_group(r[3]),"gender":_gender(r[4]),"symptoms":_split_items(r[5]),"medications":_split_items(r[6]),"duration":str(r[7] or ""),"severity":r[8],"risk":_risk(r[9])}
            rows.append(item)
    finally: conn.close()
    ag=filters.get("age_group") or ""; ge=filters.get("gender") or ""; med=_norm(filters.get("medication") or ""); sym=_norm(filters.get("symptom") or ""); risk=filters.get("risk") or ""
    out=[]
    for r in rows:
        if ag and r["age_group"]!=ag: continue
        if ge and r["gender"]!=ge: continue
        if med and not any(med in _norm(x) for x in r["medications"]): continue
        if sym and not any(sym in _norm(x) for x in r["symptoms"]): continue
        if risk and r["risk"]!=risk: continue
        out.append(r)
    return out


def _privacy_threshold() -> int:
    try: return max(3,min(20,int(os.environ.get("ANALYTICS_PRIVACY_THRESHOLD","5"))))
    except Exception: logging.getLogger(__name__).debug("Handled exception in _privacy_threshold; fallback applied (handler 268)"); return 5


def _medication_analytics(filters: dict) -> dict:
    rows=_records_for_filters(filters); k=_privacy_threshold(); users_all={r["user_hash"] for r in rows}
    # If the selected cohort is itself smaller than the privacy threshold, suppress
    # the whole report rather than leaking counts through filters or data-quality rows.
    if rows and len(users_all) < k:
        return {"rows":[],"threshold":k,"overview":[],"medication_patterns":[],"associations":[],"age_medication":[],"trends":[],"symptom_medication":[],"risk_medication":[],"data_quality":[["Privacy","Suppressed — insufficient group size"]],"preview":{"total_medication_records":None,"unique_medications":None,"most_frequent_medication":None,"most_represented_age_group":None,"top_association":None,"sufficient":False}}
    # display form for normalized medication
    display=Counter()
    for r in rows:
        for m in r["medications"]: display[(_norm(m),m)] += 1
    display_name={}
    by_norm=defaultdict(Counter)
    for (n,m),cnt in display.items(): by_norm[n][m]+=cnt
    for n,cnt in by_norm.items(): display_name[n]=cnt.most_common(1)[0][0]

    # User overview
    overview=[]
    groups=defaultdict(list)
    for r in rows: groups[(r["age_group"],r["gender"])].append(r)
    age_order=["Under 18","18–25","26–35","36–45","46–55","56+","Unknown"]
    for (ag,ge),rr in sorted(groups.items(), key=lambda x:(age_order.index(x[0][0]) if x[0][0] in age_order else 99,x[0][1])):
        users={x["user_hash"] for x in rr}
        if len(users)<k:
            overview.append([ag,ge,"Suppressed — insufficient group size","","",""])
            continue
        sc=Counter(s for x in rr for s in x["symptoms"]); rc=Counter(x["risk"] for x in rr)
        overview.append([ag,ge,len(users),len(rr),sc.most_common(1)[0][0] if sc else "—",rc.most_common(1)[0][0] if rc else "—"])

    # medication patterns & age rankings
    age_users=defaultdict(set)
    for r in rows: age_users[r["age_group"]].add(r["user_hash"])
    med_age=defaultdict(list)
    for r in rows:
        for n in set(_norm(m) for m in r["medications"] if _norm(m)): med_age[(n,r["age_group"])].append(r)
    medication_patterns=[]; age_med=[]; rank_buckets=defaultdict(list)
    for (n,ag),rr in med_age.items():
        users={x["user_hash"] for x in rr}
        if len(users)<k: continue
        sc=Counter(s for x in rr for s in x["symptoms"]); rc=Counter(x["risk"] for x in rr); denom=max(1,len(age_users[ag])); pct=len(users)/denom
        name=display_name.get(n,n)
        medication_patterns.append([name,ag,len(users),pct,sc.most_common(1)[0][0] if sc else "—",rc.most_common(1)[0][0] if rc else "—"])
        rank_buckets[ag].append((len(users),name,pct))
    for ag, vals in rank_buckets.items():
        for rank,(nusers,name,pct) in enumerate(sorted(vals,reverse=True),1): age_med.append([ag,name,nusers,pct,rank])

    # Association rules per analysis record, privacy threshold on distinct users with pair.
    records_with_med=[r for r in rows if r["medications"]]; total_med_records=len(records_with_med)
    med_user=defaultdict(set); pair_user=defaultdict(set); med_records=Counter(); pair_records=Counter()
    for r in records_with_med:
        meds=sorted(set(_norm(m) for m in r["medications"] if _norm(m)))
        for m in meds:
            med_user[m].add(r["user_hash"]); med_records[m]+=1
        for i,a in enumerate(meds):
            for b in meds[i+1:]:
                pair_user[(a,b)].add(r["user_hash"]); pair_records[(a,b)]+=1
    # Association-rule metrics use analysis records as baskets (the standard
    # support/confidence/lift definition). Distinct-user counts are used only
    # as a privacy gate so small cohorts never appear in the export.
    denom_records=max(1,total_med_records)
    associations=[]
    for (a,b),us in pair_user.items():
        if len(us)<k or len(med_user[a])<k or len(med_user[b])<k: continue
        pair_n=pair_records[(a,b)]; support=pair_n/denom_records
        for left,right in ((a,b),(b,a)):
            left_n=med_records[left]; right_support=med_records[right]/denom_records
            confidence=pair_n/max(1,left_n); lift=(confidence/right_support) if right_support else None
            associations.append([display_name.get(left,left),display_name.get(right,right),support,confidence,lift,pair_n])
    associations.sort(key=lambda x:((x[4] or 0),x[3],x[2]),reverse=True)

    # Trends monthly. Apply privacy threshold to each medication/month group.
    month_med=defaultdict(list)
    for r in rows:
        month=r["timestamp"][:7]
        for n in set(_norm(m) for m in r["medications"] if _norm(m)): month_med[(month,n)].append(r)
    trends=[]; med_month_counts=defaultdict(dict)
    for (month,n),rr in month_med.items():
        if len({x["user_hash"] for x in rr})<k: continue
        med_month_counts[n][month]=len(rr)
    for n, mp in med_month_counts.items():
        prev=None
        for month in sorted(mp):
            count=mp[month]; change=None if prev in (None,0) else (count-prev)/prev
            trends.append([month,display_name.get(n,n),count,change]); prev=count

    # Symptom/medication patterns
    sm=defaultdict(list); symptom_age_den=Counter()
    for r in rows:
        syms=set(_norm(s) for s in r["symptoms"] if _norm(s)); meds=set(_norm(m) for m in r["medications"] if _norm(m))
        for s in syms: symptom_age_den[(s,r["age_group"])]+=1
        for s in syms:
            for m in meds: sm[(s,m,r["age_group"])].append(r)
    symptom_med=[]
    for (s,m,ag),rr in sm.items():
        if len({x["user_hash"] for x in rr})<k: continue
        pct=len(rr)/max(1,symptom_age_den[(s,ag)])
        sym_disp=next((x for r in rr for x in r["symptoms"] if _norm(x)==s),s)
        symptom_med.append([sym_disp,display_name.get(m,m),len(rr),pct,ag])

    # Risk & medication
    risk_med=[]
    for n, us in med_user.items():
        if len(us)<k: continue
        relevant=[r for r in rows if any(_norm(m)==n for m in r["medications"])]
        rc=Counter(r["risk"] for r in relevant)
        risk_med.append([display_name.get(n,n),rc["Low Risk"],rc["Needs Follow-up"],rc["Urgent"],len(relevant)])
    risk_med.sort(key=lambda x:x[-1],reverse=True)

    # Data quality
    total=len(rows); missing_med=sum(not r["medications"] for r in rows); missing_age=sum(r["age_group"]=="Unknown" for r in rows); missing_gender=sum(r["gender"]=="Unknown" for r in rows)
    invalid=sum((r["age"] is not None and r["age_group"]=="Unknown") for r in rows)
    seen=set(); duplicates=0
    for r in rows:
        key=(r["user_hash"],r["timestamp"],tuple(sorted(_norm(s) for s in r["symptoms"])),tuple(sorted(_norm(m) for m in r["medications"])))
        if key in seen: duplicates+=1
        else: seen.add(key)
    # Unknown medication is only meaningful if a verified medication knowledge table exists.
    unknown_med=None
    conn=db._conn(); c=conn.cursor()
    try:
        if db.USE_POSTGRES: c.execute("SELECT 1 FROM information_schema.tables WHERE table_name=%s",("mk_medications",))
        else: c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",("mk_medications",))
        has_mk=bool(c.fetchone())
        if has_mk:
            cols=_columns(c,"mk_medications"); name_cols=[x for x in ("name","name_en","name_ar","medication_name") if x in cols]
            known=set()
            for col in name_cols:
                c.execute(f"SELECT {col} FROM mk_medications")
                known.update(_norm(r[0]) for r in c.fetchall() if r[0])
            unique=set(_norm(m) for r in rows for m in r["medications"] if _norm(m)); unknown_med=sum(1 for m in unique if m not in known)
    finally: conn.close()
    issue_cells=missing_med+missing_age+missing_gender+invalid
    score=round(max(0.0,100.0-(issue_cells/(max(1,total)*4))*100.0),1) if total else None
    excluded=sum(1 for r in rows if not r["medications"])
    dq=[["Total Records",total],["Missing Medication",missing_med],["Missing Age",missing_age],["Missing Gender",missing_gender],["Invalid Values",invalid],["Duplicate Records",duplicates],["Unknown Medications",unknown_med if unknown_med is not None else "Not available — no verified medication KB table"],["Records excluded from medication analysis",excluded],["Data Quality Score",score/100 if score is not None else None]]

    # Preview, only privacy-safe top values.
    med_counts=Counter()
    for n,us in med_user.items():
        if len(us)>=k: med_counts[display_name.get(n,n)]=len(us)
    age_counts=Counter()
    for ag,us in age_users.items():
        if len(us)>=k: age_counts[ag]=len(us)
    return {"rows":rows,"threshold":k,"overview":overview,"medication_patterns":medication_patterns,"associations":associations,"age_medication":age_med,"trends":trends,"symptom_medication":symptom_med,"risk_medication":risk_med,"data_quality":dq,"preview":{"total_medication_records":total_med_records,"unique_medications":len({n for n,us in med_user.items() if len(us)>=k}),"most_frequent_medication":med_counts.most_common(1)[0][0] if med_counts else None,"most_represented_age_group":age_counts.most_common(1)[0][0] if age_counts else None,"top_association":(associations[0][0]+" → "+associations[0][1]) if associations else None,"sufficient":bool(med_counts)}}


def medication_analytics_preview(filters: dict) -> dict:
    data=_medication_analytics(filters)
    return {"preview":data["preview"],"privacy_threshold":data["threshold"],"record_count":len(data["rows"])}


def export_medication_analytics_excel(filters: dict) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    data=_medication_analytics(filters); k=data["threshold"]
    wb=Workbook(); wb.remove(wb.active)
    navy="163B5C"; pale="EAF5FC"; line="DCE8F0"; gray="607487"
    generated=datetime.now(ZoneInfo("Asia/Riyadh")).strftime("%Y-%m-%d %H:%M %Z")
    period_label=filters.get("period") or "30d"
    if period_label=="custom": period_label=f"{filters.get('start') or '—'} to {filters.get('end') or '—'}"
    thin=Side(style="thin",color=line)

    def sheet(name,title,headers,rows,percent_cols=()):
        ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(1,len(headers)))
        ws.cell(1,1,title); ws.cell(1,1).font=Font(bold=True,size=16,color=navy); ws.cell(1,1).fill=PatternFill("solid",fgColor=pale)
        ws.cell(2,1,f"Generated: {generated}   |   Data Period: {period_label}   |   Privacy threshold: {k} anonymous users")
        ws.cell(2,1).font=Font(size=9,color=gray)
        hr=4
        for j,h in enumerate(headers,1):
            cell=ws.cell(hr,j,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        for r,row in enumerate(rows,hr+1):
            for j,val in enumerate(row,1):
                c=ws.cell(r,j,val); c.alignment=Alignment(vertical="top",wrap_text=True); c.border=Border(bottom=thin)
                if j in percent_cols and isinstance(val,(int,float)): c.number_format="0.0%"
        ws.freeze_panes=f"A{hr+1}"; ws.auto_filter.ref=f"A{hr}:{get_column_letter(len(headers))}{max(hr,ws.max_row)}"
        for col in range(1,len(headers)+1):
            vals=[str(ws.cell(r,col).value or "") for r in range(1,min(ws.max_row,200)+1)]
            ws.column_dimensions[get_column_letter(col)].width=min(38,max(12,max((len(x) for x in vals),default=10)+2))
        return ws

    ws=sheet("User Overview","SymptoSense Data Analytics Report — User Overview",["Age Group","Gender","Number of Users","Number of Analyses","Most Reported Symptom","Most Common Risk Level"],data["overview"])
    ws.insert_rows(3,2)
    ws["A3"]="Privacy: Aggregated & Anonymized — groups smaller than the privacy threshold are suppressed."
    ws["A4"]="Important: These are descriptive usage patterns and are not medical recommendations or evidence of medication effectiveness, safety, or suitability."
    # Re-create header formatting after insert shifts header to row 6.
    ws.freeze_panes="A7"; ws.auto_filter.ref=f"A6:F{ws.max_row}"

    sheet("Medication Patterns","Medication Patterns",["Medication","Age Group","Number of Users","Percentage of Users","Most Common Reported Symptom","Most Common Risk Level"],data["medication_patterns"],(4,))
    sheet("Medication Associations","Medication Associations — Association ≠ Medical Recommendation",["Medication A","Medication B","Support","Confidence","Lift","Number of Records"],data["associations"],(3,4))
    sheet("Age Medication Analysis","Age & Medication Analysis",["Age Group","Medication","Users","Percentage","Rank"],data["age_medication"],(4,))
    tws=sheet("Medication Trends","Medication Trends",["Month","Medication","Number of Records","Percentage Change"],data["trends"],(4,))
    sheet("Symptom Medication Patterns","Symptom & Medication Patterns — co-occurrence only",["Symptom","Medication","Number of Records","Percentage","Age Group"],data["symptom_medication"],(4,))
    rws=sheet("Risk Level & Medication","Risk Level & Medication — no causal inference",["Medication","Low Risk","Needs Follow-up","Urgent","Total Records"],data["risk_medication"])
    dws=sheet("Data Quality","Data Quality",["Metric","Value"],data["data_quality"])
    dws["A3"]="Data Quality Score = completeness of medication, age, gender and basic validity fields in the selected analysis records."
    for _r in range(5, dws.max_row + 1):
        if str(dws.cell(_r,1).value or "") == "Data Quality Score" and isinstance(dws.cell(_r,2).value,(int,float)):
            dws.cell(_r,2).number_format="0.0%"

    if data["trends"]:
        monthly=Counter()
        for month,med,count,change in data["trends"]: monthly[month]+=int(count or 0)
        start=tws.max_row+3; tws.cell(start,1,"Month"); tws.cell(start,2,"Total Medication Records")
        for i,(m,n) in enumerate(sorted(monthly.items()),start+1): tws.cell(i,1,m);tws.cell(i,2,n)
        if len(monthly)>=2:
            chart=LineChart(); chart.title="Medication Usage Over Time"; chart.y_axis.title="Records"; chart.x_axis.title="Month"; chart.height=7; chart.width=13
            chart.add_data(Reference(tws,min_col=2,min_row=start,max_row=start+len(monthly)),titles_from_data=True); chart.set_categories(Reference(tws,min_col=1,min_row=start+1,max_row=start+len(monthly))); tws.add_chart(chart,"F4")
    else:
        tws["A5"]="Insufficient historical data for trend analysis."

    if data["risk_medication"]:
        top=min(10,len(data["risk_medication"])); chart=BarChart(); chart.type="bar"; chart.grouping="stacked"; chart.overlap=100; chart.title="Risk Level Distribution by Medication"; chart.height=8; chart.width=14
        chart.add_data(Reference(rws,min_col=2,max_col=4,min_row=4,max_row=4+top),titles_from_data=True); chart.set_categories(Reference(rws,min_col=1,min_row=5,max_row=4+top)); rws.add_chart(chart,"G4")

    out=io.BytesIO(); wb.save(out); out.seek(0); return out
