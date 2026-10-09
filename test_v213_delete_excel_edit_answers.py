import source_bundle
from pathlib import Path
import json, os, tempfile

ROOT=Path(__file__).resolve().parent
WEB=source_bundle.webapp_text()
CHAT=source_bundle.chat_view_text()
MED=(ROOT/'medication_email.py').read_text(encoding='utf-8')
ADV=(ROOT/'advanced_features.py').read_text(encoding='utf-8')
CORE=(ROOT/'analysis_core.py').read_text(encoding='utf-8')


def test_delete_is_persistent_and_server_verified_static():
    assert 'def delete_plan(user_id, plan_id):' in MED
    assert 'DELETE FROM med_plans' in MED
    assert 'DELETE FROM medication_reminders' in MED
    assert 'medication_delete_not_persisted' in MED
    assert 'medication_email.delete_plan(uid, plan_id)' in WEB
    assert 'delete_not_persisted' in WEB
    assert "?verify_delete=" in WEB
    assert '✓ تم حذف التذكير نهائيًا' in WEB


def test_delete_persists_across_fresh_sqlite_connection(monkeypatch):
    import db, medication_email
    tmp=tempfile.NamedTemporaryFile(suffix='.db',delete=False)
    tmp.close()
    try:
        monkeypatch.setattr(db,'DATABASE_URL','')
        monkeypatch.setattr(db,'USE_POSTGRES',False)
        monkeypatch.setattr(db,'PH','?')
        monkeypatch.setattr(db,'DB_PATH',tmp.name)
        monkeypatch.setattr(db,'_DB_READY_KEY',None)
        monkeypatch.setattr(medication_email,'PH','?')
        db.init_db()
        uid=1
        plan_id=medication_email.save_plan(uid,{
            'med_name':'QA V213 DELETE','times':['08:00'],'frequency':'daily',
            'start_date':'2027-12-31','timezone':'Asia/Riyadh','delivery_channel':'email'
        })
        assert any(int(x['id'])==int(plan_id) for x in medication_email.list_plans(uid,active_only=False))
        assert medication_email.delete_plan(uid,plan_id) is True
        assert not any(int(x['id'])==int(plan_id) for x in medication_email.list_plans(uid,active_only=False))
        conn=db._conn(); c=conn.cursor(); c.execute('SELECT COUNT(*) FROM med_plans WHERE id=?',(int(plan_id),)); assert c.fetchone()[0]==0; conn.close()
    finally:
        try: os.unlink(tmp.name)
        except OSError: pass


def test_history_has_edit_answers_entry_points_and_editor_prefills():
    assert "&edit=1&from=history" in WEB
    assert "تعديل إجاباتي" in WEB
    assert "renderAnswerEditor(a.id,base);" in CHAT
    assert "conditions:a.conditions||snap.conditions||''" in CHAT
    assert "medications:a.medications||snap.medications||''" in CHAT
    assert "allergies:a.allergies||snap.allergies||''" in CHAT
    assert "bodyEl.hidden=false" in CHAT
    assert "host.scrollIntoView" in CHAT
    for field in ('editAnswerAge','editAnswerGender','editAnswerSymptoms','editAnswerDuration','editAnswerSeverity','editAnswerConditions','editAnswerMeds','editAnswerAllergies','editAnswerNotes'):
        assert field in CHAT


def test_analysis_api_rows_include_editable_context():
    assert 'r.conditions,r.medications,res.data' in ADV
    assert '"conditions": row[9]' in ADV
    assert '"medications": row[10]' in ADV
    assert '"allergies": snapshot.get("allergies")' in ADV
    assert '"notes": snapshot.get("notes")' in ADV
    for key in ('"conditions": d.get("conditions")','"medications": d.get("medications")','"allergies": d.get("allergies")','"notes": d.get("notes")'):
        assert key in CORE


def test_admin_preview_matches_excel_source():
    import db
    from openpyxl import load_workbook
    tmp=tempfile.NamedTemporaryFile(suffix='.db',delete=False)
    tmp.close()
    try:
        db.DATABASE_URL=''; db.USE_POSTGRES=False; db.PH='?'; db.DB_PATH=tmp.name; db._DB_READY_KEY=None
        db.init_db()
        db.save_analysis_record_with_result(
            'account-v213','ar',31,'f',['غثيان','ألم بطن'],'1-3 أيام',2,'low',
            {'ok':True,'input_snapshot':{'symptoms':['غثيان','ألم بطن']},'symptoms':['غثيان','ألم بطن']}
        )
        rows=db.admin_symptom_trial_rows()
        assert rows and 'غثيان' in rows[0]['symptoms'] and 'ألم بطن' in rows[0]['symptoms']
        wb=load_workbook(db.admin_export_symptom_trials_xlsx(),data_only=True)
        ws=wb['Symptom Trials']
        assert ws['C4'].value=='Symptoms | الأعراض'
        exported='\n'.join(str(ws.cell(r,3).value or '') for r in range(5,ws.max_row+1))
        assert 'غثيان' in exported and 'ألم بطن' in exported
        assert 'معاينة بيانات تحليل الأعراض' in WEB
        assert '__SYMPTOM_PREVIEW__' in WEB
    finally:
        try: os.unlink(tmp.name)
        except OSError: pass


def test_localized_admin_path_redirects_to_canonical_admin():
    assert 'if subpath == "admin" or subpath.startswith("admin/")' in WEB
    assert 'return redirect("/" + subpath, code=302)' in WEB
